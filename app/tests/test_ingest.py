from pathlib import Path

import pytest

from rag.ingest.chunker import chunk_text
from rag.ingest.loaders import LoaderOptions, UnsupportedFormat, extract, html_to_text


def test_chunker_respects_size_and_overlap():
    text = "\n\n".join(f"Paragraph {i}. " + "word " * 80 for i in range(20))
    chunks = chunk_text(text, size=1000, overlap=100)
    assert len(chunks) > 1
    assert all(len(c) <= 1000 + 100 + 1 for c in chunks)
    assert "Paragraph 0." in chunks[0]
    assert "Paragraph 19." in chunks[-1]


def test_chunker_short_text_single_chunk():
    assert chunk_text("hello world", size=100, overlap=10) == ["hello world"]
    assert chunk_text("   ", size=100) == []


def test_html_to_text_keeps_structure():
    html = "<h2>Title</h2><p>One</p><table><tr><th>A</th><th>B</th></tr><tr><td>1</td><td>2</td></tr></table><ul><li>x</li></ul>"
    text = html_to_text(html)
    assert "## Title" in text
    assert "A | B" in text and "1 | 2" in text
    assert "- x" in text


def test_text_and_unknown_extension(tmp_path: Path):
    f = tmp_path / "config.conf"
    f.write_text("key=value\nserver=internal", encoding="utf-8")
    assert "server=internal" in extract(f)[0].text


def test_csv(tmp_path: Path):
    f = tmp_path / "data.csv"
    f.write_text("name;port\nqdrant;6333\nollama;11434\n", encoding="utf-8")
    text = extract(f)[0].text
    assert "name: qdrant; port: 6333" in text


def test_docx(tmp_path: Path):
    import docx

    d = docx.Document()
    d.add_heading("Procedure", level=1)
    d.add_paragraph("Step one of the deployment")
    table = d.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Environment", "URL"
    table.cell(1, 0).text, table.cell(1, 1).text = "QA", "qa.local"
    path = tmp_path / "doc.docx"
    d.save(path)
    text = extract(path)[0].text
    assert "# Procedure" in text
    assert "Step one of the deployment" in text
    assert "Environment: QA; URL: qa.local" in text


def test_locations_follow_the_language(tmp_path: Path):
    import openpyxl

    wb = openpyxl.Workbook()
    wb.active.title = "Sales"
    wb.active.append(["a", "b"])
    path = tmp_path / "book.xlsx"
    wb.save(path)
    assert extract(path)[0].location == "sheet Sales"
    assert extract(path, LoaderOptions(language="es"))[0].location == "hoja Sales"


def test_binary_is_unsupported(tmp_path: Path):
    f = tmp_path / "file.xyz"
    f.write_bytes(bytes(range(256)) * 10)
    with pytest.raises(UnsupportedFormat):
        extract(f)


def test_image_without_ocr_is_unsupported(tmp_path: Path):
    f = tmp_path / "screenshot.png"
    f.write_bytes(b"\x89PNG\r\n\x1a\n")
    with pytest.raises(UnsupportedFormat):
        extract(f)
