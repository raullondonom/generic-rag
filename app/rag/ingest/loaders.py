"""Text extraction from any readable format.

Each loader returns a list of `Section` (text + location inside the file, e.g. "page 3" or
"sheet Sales", in the configured language). Unknown binary files raise `UnsupportedFormat`;
read errors propagate so they can be recorded.
"""

from __future__ import annotations

import csv
import email
import email.policy
import io
import re
import shutil
import subprocess
from dataclasses import dataclass
from pathlib import Path
from typing import Callable

from bs4 import BeautifulSoup
from charset_normalizer import from_bytes

from ..i18n import t


@dataclass
class Section:
    text: str
    location: str = ""


class UnsupportedFormat(Exception):
    pass


@dataclass
class LoaderOptions:
    ocr_enabled: bool = False
    ocr_langs: str = "eng+spa"
    language: str = "en"


# ---------------------------------------------------------------- helpers

def decode_bytes(data: bytes) -> str:
    if not data:
        return ""
    for bom, enc in ((b"\xef\xbb\xbf", "utf-8-sig"), (b"\xff\xfe", "utf-16"), (b"\xfe\xff", "utf-16")):
        if data.startswith(bom):
            return data.decode(enc, errors="replace")
    try:
        return data.decode("utf-8")
    except UnicodeDecodeError:
        best = from_bytes(data).best()
        return str(best) if best else data.decode("latin-1", errors="replace")


def looks_binary(data: bytes) -> bool:
    sample = data[:8192]
    if b"\x00" in sample:
        return True
    if not sample:
        return False
    control = sum(1 for b in sample if b < 9 or 13 < b < 32)
    return control / len(sample) > 0.05


def clean_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\xa0", " ")
    text = "\n".join(line.rstrip() for line in text.split("\n"))
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def html_to_text(html: str) -> str:
    """HTML -> text keeping headings (#), tables (cells with |) and lists (-)."""
    soup = BeautifulSoup(html, "lxml")
    for tag in soup(["script", "style", "noscript"]):
        tag.decompose()
    # Inside out so nested tables or lists are not lost
    for tr in reversed(soup.find_all("tr")):
        cells = [c.get_text(" ", strip=True) for c in tr.find_all(["td", "th"], recursive=False)]
        tr.replace_with(soup.new_string("\n" + " | ".join(cells) + "\n"))
    for li in reversed(soup.find_all("li")):
        li.replace_with(soup.new_string("\n- " + li.get_text(" ", strip=True) + "\n"))
    for level in range(1, 7):
        for h in soup.find_all(f"h{level}"):
            h.replace_with(soup.new_string(f"\n\n{'#' * level} {h.get_text(' ', strip=True)}\n\n"))
    for br in soup.find_all("br"):
        br.replace_with(soup.new_string("\n"))
    for block in soup.find_all(["p", "div", "pre", "section", "article", "blockquote"]):
        block.insert_before(soup.new_string("\n"))
        block.insert_after(soup.new_string("\n"))
    return clean_text(soup.get_text())


def table_rows_to_text(rows: list[list[str]]) -> str:
    """Table rows -> 'Header: value; ...' text (the first non-empty row is the header)."""
    rows = [[str(c).strip() if c is not None else "" for c in r] for r in rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return ""
    header = rows[0]
    if len(rows) == 1 or not all(header):
        return "\n".join(" | ".join(r) for r in rows)
    lines = [" | ".join(header)]
    for r in rows[1:]:
        pairs = [f"{header[i] if i < len(header) else f'col{i + 1}'}: {v}" for i, v in enumerate(r) if v]
        lines.append("; ".join(pairs))
    return "\n".join(lines)


def ocr_image(image, langs: str) -> str:
    import pytesseract

    return pytesseract.image_to_string(image, lang=langs)


# ---------------------------------------------------------------- loaders

def load_pdf(path: Path, opts: LoaderOptions) -> list[Section]:
    import fitz  # pymupdf

    sections = []
    with fitz.open(path) as doc:
        for i, page in enumerate(doc, start=1):
            text = page.get_text("text")
            if not text.strip() and opts.ocr_enabled:
                from PIL import Image

                pix = page.get_pixmap(dpi=200)
                img = Image.open(io.BytesIO(pix.tobytes("png")))
                text = ocr_image(img, opts.ocr_langs)
            if text.strip():
                sections.append(Section(clean_text(text), t(opts.language, "page", n=i)))
    return sections


def load_docx(path: Path, opts: LoaderOptions) -> list[Section]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(str(path))
    parts: list[str] = []
    for child in document.element.body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, document)
            text = p.text.strip()
            if not text:
                continue
            style = (p.style.name or "").lower() if p.style is not None else ""  # localized in non-English Word
            m = re.search(r"(heading|título|titulo)\s*(\d)", style)
            if m:
                parts.append(f"\n{'#' * int(m.group(2))} {text}\n")
            elif style in ("title", "título", "titulo"):
                parts.append(f"\n# {text}\n")
            else:
                parts.append(text)
        elif tag == "tbl":
            table = Table(child, document)
            rows = [[cell.text for cell in row.cells] for row in table.rows]
            parts.append(table_rows_to_text(rows))
    return [Section(clean_text("\n".join(parts)))]


def load_doc(path: Path, opts: LoaderOptions) -> list[Section]:
    if not shutil.which("antiword"):
        raise UnsupportedFormat(t(opts.language, "doc_needs_antiword"))
    out = subprocess.run(["antiword", str(path)], capture_output=True, timeout=120)
    if out.returncode != 0:
        raise RuntimeError(out.stderr.decode(errors="replace")[:300])
    return [Section(clean_text(decode_bytes(out.stdout)))]


def load_pptx(path: Path, opts: LoaderOptions) -> list[Section]:
    from pptx import Presentation

    prs = Presentation(str(path))
    sections = []
    for i, slide in enumerate(prs.slides, start=1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [[c.text for c in r.cells] for r in shape.table.rows]
                parts.append(table_rows_to_text(rows))
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append(t(opts.language, "notes") + ": " + notes)
        text = clean_text("\n".join(parts))
        if text:
            sections.append(Section(text, t(opts.language, "slide", n=i)))
    return sections


def load_xlsx(path: Path, opts: LoaderOptions) -> list[Section]:
    import openpyxl

    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    sections = []
    try:
        for ws in wb.worksheets:
            rows = [list(r) for r in ws.iter_rows(values_only=True)]
            text = table_rows_to_text(rows)
            if text:
                sections.append(Section(text, t(opts.language, "sheet", name=ws.title)))
    finally:
        wb.close()
    return sections


def load_xls(path: Path, opts: LoaderOptions) -> list[Section]:
    import xlrd

    book = xlrd.open_workbook(str(path))
    sections = []
    for sheet in book.sheets():
        rows = [sheet.row_values(i) for i in range(sheet.nrows)]
        text = table_rows_to_text(rows)
        if text:
            sections.append(Section(text, t(opts.language, "sheet", name=sheet.name)))
    return sections


def load_csv(path: Path, opts: LoaderOptions) -> list[Section]:
    text = decode_bytes(path.read_bytes())
    delimiter = "\t" if path.suffix.lower() == ".tsv" else None
    if delimiter is None:
        try:
            delimiter = csv.Sniffer().sniff(text[:5000], delimiters=",;|\t").delimiter
        except csv.Error:
            delimiter = ","
    rows = list(csv.reader(io.StringIO(text), delimiter=delimiter))
    return [Section(table_rows_to_text(rows))]


def load_odf(path: Path, opts: LoaderOptions) -> list[Section]:
    from odf import teletype
    from odf.namespaces import TEXTNS
    from odf.opendocument import load as odf_load

    doc = odf_load(str(path))
    parts: list[str] = []

    def walk(node) -> None:
        for child in node.childNodes:
            qname = getattr(child, "qname", None)
            if qname == (TEXTNS, "h"):
                parts.append(f"\n# {teletype.extractText(child)}\n")
            elif qname == (TEXTNS, "p"):
                parts.append(teletype.extractText(child))
            else:
                walk(child)

    walk(doc.body)
    return [Section(clean_text("\n".join(parts)))]


def load_rtf(path: Path, opts: LoaderOptions) -> list[Section]:
    from striprtf.striprtf import rtf_to_text

    return [Section(clean_text(rtf_to_text(decode_bytes(path.read_bytes()), errors="ignore")))]


def load_epub(path: Path, opts: LoaderOptions) -> list[Section]:
    import ebooklib
    from ebooklib import epub

    book = epub.read_epub(str(path))
    sections = []
    for i, item in enumerate(book.get_items_of_type(ebooklib.ITEM_DOCUMENT), start=1):
        text = html_to_text(item.get_content().decode("utf-8", errors="replace"))
        if text:
            sections.append(Section(text, t(opts.language, "chapter", n=i)))
    return sections


def load_html(path: Path, opts: LoaderOptions) -> list[Section]:
    return [Section(html_to_text(decode_bytes(path.read_bytes())))]


def load_eml(path: Path, opts: LoaderOptions) -> list[Section]:
    msg = email.message_from_bytes(path.read_bytes(), policy=email.policy.default)
    header = t(opts.language, "email_header", subject=msg["subject"], sender=msg["from"], to=msg["to"], date=msg["date"])
    body = msg.get_body(preferencelist=("plain", "html"))
    content = ""
    if body is not None:
        content = body.get_content()
        if body.get_content_type() == "text/html":
            content = html_to_text(content)
    return [Section(clean_text(header + "\n" + content))]


def load_msg(path: Path, opts: LoaderOptions) -> list[Section]:
    import extract_msg

    msg = extract_msg.Message(str(path))
    try:
        header = t(opts.language, "email_header", subject=msg.subject, sender=msg.sender, to=msg.to, date=msg.date)
        body = msg.body or (html_to_text(msg.htmlBody.decode(errors="replace")) if msg.htmlBody else "")
    finally:
        msg.close()
    return [Section(clean_text(header + "\n" + body))]


def load_image(path: Path, opts: LoaderOptions) -> list[Section]:
    if not opts.ocr_enabled:
        raise UnsupportedFormat(t(opts.language, "image_no_ocr"))
    from PIL import Image

    with Image.open(path) as img:
        return [Section(clean_text(ocr_image(img, opts.ocr_langs)))]


def load_text(path: Path, opts: LoaderOptions) -> list[Section]:
    data = path.read_bytes()
    if looks_binary(data):
        raise UnsupportedFormat(t(opts.language, "binary"))
    return [Section(clean_text(decode_bytes(data)))]


LOADERS: dict[str, Callable[[Path, LoaderOptions], list[Section]]] = {
    ".pdf": load_pdf,
    ".docx": load_docx,
    ".docm": load_docx,
    ".doc": load_doc,
    ".pptx": load_pptx,
    ".pptm": load_pptx,
    ".xlsx": load_xlsx,
    ".xlsm": load_xlsx,
    ".xls": load_xls,
    ".csv": load_csv,
    ".tsv": load_csv,
    ".odt": load_odf,
    ".ods": load_odf,
    ".odp": load_odf,
    ".rtf": load_rtf,
    ".epub": load_epub,
    ".html": load_html,
    ".htm": load_html,
    ".xhtml": load_html,
    ".eml": load_eml,
    ".msg": load_msg,
    ".png": load_image,
    ".jpg": load_image,
    ".jpeg": load_image,
    ".tif": load_image,
    ".tiff": load_image,
    ".bmp": load_image,
    ".gif": load_image,
    ".webp": load_image,
}

# Known binaries that make no sense to read as text
SKIP_EXTENSIONS = {
    ".zip", ".7z", ".rar", ".tar", ".gz", ".exe", ".dll", ".bin", ".iso", ".mp3", ".mp4",
    ".avi", ".mov", ".mkv", ".wav", ".class", ".jar", ".pyc", ".so", ".db", ".sqlite",
    ".svg", ".ico", ".woff", ".woff2", ".ttf", ".otf",
}


def extract(path: Path, opts: LoaderOptions | None = None) -> list[Section]:
    """Extracts the text of a file. Unknown extensions are read as text."""
    opts = opts or LoaderOptions()
    ext = path.suffix.lower()
    if ext in SKIP_EXTENSIONS:
        raise UnsupportedFormat(t(opts.language, "unsupported_ext", ext=ext))
    loader = LOADERS.get(ext, load_text)
    sections = [s for s in loader(path, opts) if s.text.strip()]
    return sections
